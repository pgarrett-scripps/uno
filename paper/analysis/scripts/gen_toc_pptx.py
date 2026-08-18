#!/usr/bin/env python3
"""Scaffold the TOC / graphical abstract as an editable PowerPoint slide.

Writes paper/toc_graphic.pptx: ONE slide, 6.5 in x 3.5 in, which is exactly
twice the ACS "For Table of Contents Only" box (3.25 in x 1.75 in). Everything
on it is a native shape or text box, so the slide can be reworked freely in
PowerPoint. Export at 50% scale (or 300 dpi over the half-size box) for
submission.

This script is a one-shot scaffold, not part of `just assets`: once the slide
is edited by hand in PowerPoint, RERUNNING THIS SCRIPT OVERWRITES THE EDITS.
It is deliberately named outside the `fig*.py` and `gen_*_table.py` globs.

Same concept as figures/fig_toc_graphic.png, drawn with the manuscript
palette (Wong): an MS1 feature map (no MS/MS), matched against an in-silico
peptide database on mass and predicted RT, out to an FDR-controlled protein
list. No data anywhere; the dot cloud is a seeded random scatter.
"""

from __future__ import annotations

import random
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
import style  # noqa: E402  (palette + repo paths only; no matplotlib output)

OUT = style.PAPER / "toc_graphic.pptx"

UNO = RGBColor(0x00, 0x72, 0xB2)
UNO_FILL = RGBColor(0xE4, 0xEF, 0xF7)
TARGET = RGBColor(0x00, 0x9E, 0x73)
TARGET_FILL = RGBColor(0xE6, 0xF4, 0xEF)
GREY = RGBColor(0x7F, 0x7F, 0x7F)
NEUTRAL = RGBColor(0x33, 0x33, 0x33)
NOTE = RGBColor(0x5A, 0x7A, 0x8C)

SLIDE_W, SLIDE_H = 6.5, 3.5


def _no_shadow(shape):
    shape.shadow.inherit = False


def add_box(shapes, x, y, w, h, fill, line, line_w=1.25, shape_type=None):
    sp = shapes.add_shape(shape_type or MSO_SHAPE.ROUNDED_RECTANGLE,
                          Inches(x), Inches(y), Inches(w), Inches(h))
    _no_shadow(sp)
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    sp.line.color.rgb = line
    sp.line.width = Pt(line_w)
    return sp


def add_text(shapes, x, y, w, h, runs, size, align=PP_ALIGN.CENTER,
             color=NEUTRAL, bold=False, italic=False, rotation=None):
    """runs: a string, or list of (text, overrides-dict) for mixed styling."""
    tb = shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    if isinstance(runs, str):
        runs = [(runs, {})]
    lines = []
    for text, ov in runs:
        for i, part in enumerate(text.split("\n")):
            if i or not lines:
                lines.append([])
            lines[-1].append((part, ov))
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for part, ov in line:
            r = p.add_run()
            r.text = part
            f = r.font
            f.name = ov.get("name", "Arial")
            f.size = Pt(ov.get("size", size))
            f.bold = ov.get("bold", bold)
            f.italic = ov.get("italic", italic)
            f.color.rgb = ov.get("color", color)
    if rotation is not None:
        tb.rotation = rotation
    return tb


def add_line(shapes, x0, y0, x1, y1, color=NEUTRAL, w=1.5, arrow=False,
             dash=None):
    conn = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x0), Inches(y0),
                                Inches(x1), Inches(y1))
    _no_shadow(conn)
    conn.line.color.rgb = color
    conn.line.width = Pt(w)
    ln = conn.line._get_or_add_ln()
    if dash:
        d = ln.makeelement(qn("a:prstDash"), {"val": dash})
        ln.append(d)
    if arrow:
        tail = ln.makeelement(qn("a:tailEnd"),
                              {"type": "triangle", "w": "med", "len": "med"})
        ln.append(tail)
    return conn


def add_dot(shapes, cx, cy, d, fill, line=None):
    sp = shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2),
                          Inches(d), Inches(d))
    _no_shadow(sp)
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(1.25)
    return sp


def main() -> int:
    prs = Presentation()
    prs.slide_width = Emu(int(Inches(SLIDE_W)))
    prs.slide_height = Emu(int(Inches(SLIDE_H)))
    shapes = prs.slides.add_slide(prs.slide_layouts[6]).shapes

    # ---------------- headline ----------------------------------------------
    add_text(shapes, 0.25, 0.08, 6.0, 0.36,
             "protein identification from MS1 features alone",
             15, bold=True)

    # ---------------- left: one run's MS1 feature map ------------------------
    px0, px1, py_top, py_bot = 0.30, 2.12, 0.80, 2.66
    add_line(shapes, px0, py_top, px0, py_bot, color=GREY, w=1.0)
    add_line(shapes, px0, py_bot, px1, py_bot, color=GREY, w=1.0)
    rng = random.Random(7)
    for _ in range(60):
        d = 0.022 + rng.random() * 0.03
        add_dot(shapes,
                px0 + 0.08 + rng.random() * (px1 - px0 - 0.16),
                py_top + 0.08 + rng.random() * (py_bot - py_top - 0.18),
                d, RGBColor(0xB8, 0xB8, 0xB8))
    # the one feature being searched: isotope envelope, circled
    hx, hy = 1.00, 2.00
    for i, sh in enumerate((0.23, 0.15, 0.085)):
        add_line(shapes, hx + i * 0.07, hy, hx + i * 0.07, hy - sh,
                 color=UNO, w=2.0)
    add_dot(shapes, hx + 0.07, hy - 0.10, 0.40, None, line=UNO)
    add_text(shapes, px0 - 0.32, (py_top + py_bot) / 2 - 0.12, 0.6, 0.24,
             "m/z", 9, color=NOTE, rotation=270)
    add_text(shapes, (px0 + px1) / 2 - 0.3, py_bot + 0.03, 0.6, 0.2,
             "RT", 9, color=NOTE)
    add_text(shapes, 0.10, 2.95, 2.2, 0.22, "one run's MS1 feature map", 10)
    add_text(shapes, 0.10, 3.17, 2.2, 0.2, "mass, RT and intensity only", 9,
             color=NOTE, italic=True)

    # "no MS/MS" mark, top-right corner of the map
    cx, cy, r = 1.80, 1.10, 0.21
    for i, sh in enumerate((0.11, 0.18, 0.09, 0.14)):
        sx = cx - 0.11 + i * 0.074
        add_line(shapes, sx, cy + 0.10, sx, cy + 0.10 - sh, color=GREY, w=1.2)
    add_dot(shapes, cx, cy, 2 * r, None, line=NEUTRAL)
    dd = r / (2 ** 0.5)
    add_line(shapes, cx - dd, cy + dd, cx + dd, cy - dd, color=NEUTRAL, w=1.25)
    add_text(shapes, cx - 0.5, cy + r + 0.02, 1.0, 0.2, "no MS/MS", 9,
             color=NOTE, italic=True)

    # ---------------- middle: the search engine ------------------------------
    bx0, bw, by0, bh = 2.42, 1.70, 1.56, 0.90
    add_text(shapes, bx0 - 0.2, 0.50, bw + 0.4, 0.2,
             "protein FASTA → digest", 9, color=NOTE, italic=True)
    add_text(shapes, bx0 - 0.2, 0.72, bw + 0.4, 0.22,
             "in-silico peptide database", 10)
    for i in range(3):
        add_box(shapes, bx0 + 0.52 + i * 0.05, 0.98 + i * 0.10, 0.62, 0.075,
                RGBColor(0xFF, 0xFF, 0xFF), GREY, 1.0)
    add_line(shapes, bx0 + bw / 2, 1.32, bx0 + bw / 2, by0 - 0.02, arrow=True)

    add_box(shapes, bx0, by0, bw, bh, UNO_FILL, UNO, 1.75)
    add_text(shapes, bx0, by0 + 0.08, bw, 0.34, "uno", 16, bold=True,
             color=UNO)
    add_text(shapes, bx0, by0 + 0.44, bw, 0.40,
             "match by mass and\npredicted RT", 10)
    add_text(shapes, bx0, by0 + bh + 0.05, bw, 0.2, "MS1 search engine", 9,
             color=NOTE, italic=True)

    add_line(shapes, px1 + 0.08, 2.0, bx0 - 0.04, 2.0, arrow=True)
    add_line(shapes, bx0 + bw + 0.04, 2.0, 4.46, 2.0, arrow=True)

    # ---------------- right: FDR-controlled protein list ---------------------
    gx0 = 4.56
    for i, frac in enumerate((1.0, 0.76, 0.53)):
        y = 1.20 + i * 0.42
        add_box(shapes, gx0, y, 1.70 * frac, 0.27, TARGET_FILL, TARGET, 1.25)
        add_text(shapes, gx0 + 0.10, y + 0.025, 0.9, 0.22, "protein", 10,
                 align=PP_ALIGN.LEFT)
    add_text(shapes, gx0 + 0.6, 2.42, 0.5, 0.2, "⋮", 10)
    add_text(shapes, 4.30, 2.95, 2.1, 0.22, "protein-level identity", 10)
    add_text(shapes, 4.30, 3.17, 2.1, 0.2, "target/decoy FDR control", 9,
             color=NOTE, italic=True)

    prs.save(OUT)
    print(f"wrote {OUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
